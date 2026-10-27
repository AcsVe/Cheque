"""Cheque register, payees, banks and accounts — every list with add / edit / multi-select bulk actions."""
from decimal import Decimal

from PySide6.QtCore import QDate, Signal
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDateEdit, QDialog, QDialogButtonBox, QHBoxLayout, QLabel,
                               QLineEdit, QPlainTextEdit, QSpinBox, QSplitter, QVBoxLayout, QWidget, QFileDialog)
from PySide6.QtCore import Qt

from ..core.db import CROSSINGS, STATUSES
from ..core.render import cheque_values
from ..core.tafqeet import CURRENCIES
from .common import DataTable, L, confirm, error, info
from .printing import print_jobs
from .theme import a11y, button, card, field, fit, fit_to_screen, form_grid, ltr, name_buttons, scroll_page


def _st(key):
    return L(*STATUSES.get(key, (key, key)))


# ───────────────────────────── register ─────────────────────────────
class RegisterTab(QWidget):
    open_cheque = Signal(int)
    changed = Signal()

    COLS = None

    def __init__(self, ctx):
        super().__init__()
        self.ctx, self.db = ctx, ctx.db
        self.d_from = ltr(QDateEdit(QDate.currentDate().addMonths(-12)), 150)
        self.d_to = ltr(QDateEdit(QDate.currentDate().addMonths(12)), 150)
        for d in (self.d_from, self.d_to):
            d.setCalendarPopup(True)
            d.setDisplayFormat('dd/MM/yyyy')
            d.dateChanged.connect(self.reload)
        self.total = QLabel()
        self.total.setProperty('chip', True)
        a11y(self.total, L('مجموع المبالغ الظاهرة', 'Total of the amounts shown'))
        self.table = DataTable([
            (L('رقم الشيك', 'Cheque no.'), 'num'), (L('التاريخ', 'Date'), 'text'), (L('المستفيد', 'Payee'), 'text'),
            (L('المبلغ', 'Amount'), 'num'), (L('العملة', 'Currency'), 'select'), (L('البنك', 'Bank'), 'select'),
            (L('الحساب', 'Account'), 'select'), (L('التسطير', 'Crossing'), 'select'), (L('الحالة', 'Status'), 'select'),
            (L('مرات الطباعة', 'Printed'), 'num'), (L('البيان', 'Memo'), 'text'),
        ], actions=[
            (L('فتح للتعديل', 'Open'), self._open, False, 'edit'),
            (L('طباعة المحدد', 'Print selected'), self._print, False, 'print'),
            (L('إلغاء الشيك', 'Cancel cheque'), self._cancel, True, 'cancel'),
            (L('حذف', 'Delete'), self._delete, True, 'trash'),
        ], export_name='cheques', title=L('سجل الشيكات', 'Cheque register'))
        self.table.activated.connect(self.open_cheque.emit)
        self.table.proxy.layoutChanged.connect(self._totals)
        self.table.proxy.rowsInserted.connect(self._totals)
        self.table.proxy.rowsRemoved.connect(self._totals)
        self.table.proxy.modelReset.connect(self._totals)
        top = QHBoxLayout()
        top.setSpacing(8)
        for text, w in ((L('من تاريخ', 'From'), self.d_from), (L('إلى تاريخ', 'To'), self.d_to)):
            lab = QLabel(text)
            lab.setBuddy(w)
            a11y(w, text)
            top.addWidget(lab)
            top.addWidget(w)
            top.addSpacing(8)
        top.addStretch(1)
        top.addWidget(self.total)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(10)
        lay.addLayout(top)
        lay.addWidget(self.table)
        self.reload()

    def reload(self, *_):
        rows = []
        for c in self.db.cheques('c.cheque_date BETWEEN ? AND ?',
                                 (self.d_from.date().toString('yyyy-MM-dd'), self.d_to.date().toString('yyyy-MM-dd'))):
            amount = Decimal(c['amount'])
            rows.append((c['id'], [
                (c['cheque_no'], int(c['cheque_no'])),
                (QDate.fromString(c['cheque_date'], 'yyyy-MM-dd').toString('dd/MM/yyyy'), c['cheque_date']),
                c['payee'], (f'{amount:,.3f}' if c['currency'] == 'JOD' else f'{amount:,.2f}', str(amount)),
                c['currency'], c['bank_name'], c['account_name'], L(*CROSSINGS.get(c['crossing'], ('', ''))),
                _st(c['status']), (c['print_count'], c['print_count']), c.get('memo') or '']))
        self.table.set_rows(rows)
        self._totals()

    def _totals(self, *_):
        total = self.table.column_sum(3)
        self.total.setText(L(f'المجموع الظاهر: {total:,.3f}', f'Total shown: {total:,.3f}'))

    def _open(self, ids):
        if ids:
            self.open_cheque.emit(ids[0])

    def _print(self, ids):
        if not self.ctx.license().can_print:
            self.ctx.license_blocked(self)
            return
        cheques = [c for c in (self.db.cheque(i) for i in ids) if c and c['status'] != 'cancelled']
        if not cheques:
            return
        if not confirm(self, L(f'طباعة {len(cheques)} من الشيكات؟ ضع الشيكات في الطابعة بالترتيب نفسه.',
                               f'Print {len(cheques)} cheque(s)? Put the cheques in the printer in the same order.')):
            return
        jobs = []
        for c in sorted(cheques, key=lambda x: (x['account_id'], int(x['cheque_no']))):
            tpl = self.db.template(c['bank_id'])
            jobs.append((tpl, cheque_values(tpl, c, self.ctx.crossing_texts())))
        if print_jobs(self, jobs, self.db.get, self.db.set):
            for c in cheques:
                self.db.mark_printed(c['id'])
        self.reload()
        self.changed.emit()

    def _cancel(self, ids):
        if confirm(self, L(f'إلغاء {len(ids)} من الشيكات؟ يبقى رقم الشيك مسجّلاً ولا يُستخدم مرة أخرى.',
                           f'Cancel {len(ids)} cheque(s)? The number stays recorded and is not reused.')):
            self.db.set_status(ids, 'cancelled')
            self.reload()
            self.changed.emit()

    def _delete(self, ids):
        if not confirm(self, L(f'حذف {len(ids)} من الشيكات نهائياً؟ الشيكات المطبوعة لا تُحذف (تُلغى بدلاً من ذلك).',
                               f'Delete {len(ids)} cheque(s)? Printed cheques are not deleted (cancel them instead).')):
            return
        done, kept = self.db.delete_cheques(ids)
        msg = L(f'حُذف: {done}', f'Deleted: {done}')
        if kept:
            msg += '\n' + L('لم تُحذف لأنها مطبوعة: ', 'Not deleted (printed): ') + '، '.join(map(str, kept))
        info(self, msg)
        self.reload()
        self.changed.emit()


# ───────────────────────────── payees ─────────────────────────────
def _dialog_layout(dlg, grid):
    """A form dialog: the fields in a card, Save (primary) and Cancel under it."""
    from PySide6.QtWidgets import QSizePolicy
    box = card()
    grid.setColumnStretch(1, 1)
    box.setLayout(grid)
    bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
    save = bb.button(QDialogButtonBox.Save)
    save.setText(L('حفظ', 'Save'))
    save.setProperty('primary', True)
    save.setMinimumWidth(110)
    cancel = bb.button(QDialogButtonBox.Cancel)
    cancel.setText(L('إلغاء', 'Cancel'))
    cancel.setMinimumWidth(110)
    name_buttons(bb)
    bb.accepted.connect(dlg.accept)
    bb.rejected.connect(dlg.reject)
    page = QWidget()
    pl = QVBoxLayout(page)
    pl.setContentsMargins(18, 8, 18, 4)
    pl.addWidget(box)
    pl.addStretch(1)
    lay = QVBoxLayout(dlg)
    lay.setContentsMargins(0, 0, 0, 18)
    lay.setSpacing(10)
    lay.addWidget(scroll_page(page), 1)
    bl = QHBoxLayout()
    bl.setContentsMargins(18, 0, 18, 0)
    bl.addWidget(bb)
    lay.addLayout(bl)
    want = page.sizeHint()
    fit_to_screen(dlg, max(560, want.width() + 4), want.height() + bb.sizeHint().height() + 32)


class PayeeDialog(QDialog):
    def __init__(self, parent, payee=None):
        super().__init__(parent)
        self.setWindowTitle(L('مستفيد', 'Payee'))
        self.name = QLineEdit((payee or {}).get('name', ''))
        self.name_en = QLineEdit((payee or {}).get('name_en') or '')
        self.phone = QLineEdit((payee or {}).get('phone') or '')
        self.notes = QLineEdit((payee or {}).get('notes') or '')
        self.active = a11y(QCheckBox(L('فعّال', 'Active')), L('المستفيد فعّال', 'Payee active'))
        self.active.setChecked(bool((payee or {}).get('is_active', 1)))
        self.name.setMinimumWidth(260)
        self.name_en.setLayoutDirection(Qt.LeftToRight)
        ltr(self.phone, 200)
        g = form_grid()
        field(g, 0, 0, L('الاسم *', 'Name *'), self.name)
        field(g, 1, 0, L('الاسم بالإنجليزية', 'English name'), self.name_en)
        field(g, 2, 0, L('الهاتف', 'Phone'), self.phone)
        field(g, 3, 0, L('ملاحظات', 'Notes'), self.notes)
        g.addWidget(self.active, 4, 1)
        _dialog_layout(self, g)

    def values(self):
        return {'name': self.name.text().strip(), 'name_en': self.name_en.text().strip() or None,
                'phone': self.phone.text().strip() or None, 'notes': self.notes.text().strip() or None,
                'is_active': 1 if self.active.isChecked() else 0}


class ImportPayeesDialog(QDialog):
    """Paste a list of payees (one per line, or name ⇥ English name ⇥ phone from Excel), or load an Excel / CSV file."""

    def __init__(self, parent):
        super().__init__(parent)
        self.setWindowTitle(L('إضافة قائمة مستفيدين', 'Add a list of payees'))
        self.box = QPlainTextEdit()
        self.box.setPlaceholderText(L('اسم في كل سطر، أو الصق من Excel: الاسم ⇥ الاسم بالإنجليزية ⇥ الهاتف',
                                      'One name per line, or paste from Excel: name ⇥ English name ⇥ phone'))
        a11y(self.box, L('قائمة المستفيدين', 'List of payees'),
             L('اسم في كل سطر، أو أعمدة: الاسم، الاسم بالإنجليزية، الهاتف', 'One name per line, or columns: name, English name, phone'))
        load = button(L('اختيار ملف Excel / CSV', 'Choose an Excel / CSV file'), 'folder')
        load.clicked.connect(self.from_file)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText(L('إضافة', 'Add'))
        bb.button(QDialogButtonBox.Ok).setProperty('primary', True)
        bb.button(QDialogButtonBox.Cancel).setText(L('إلغاء', 'Cancel'))
        name_buttons(bb)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        row = QHBoxLayout()
        row.addWidget(load)
        row.addStretch(1)
        row.addWidget(bb)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(18, 18, 18, 18)
        lay.setSpacing(12)
        lay.addWidget(self.box)
        lay.addLayout(row)
        fit_to_screen(self, 600, 400)

    def from_file(self):
        path, _ = QFileDialog.getOpenFileName(self, '', '', 'Excel / CSV (*.xlsx *.csv *.txt)')
        if not path:
            return
        lines = []
        if path.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
            for r in ws.iter_rows(values_only=True):
                lines.append('\t'.join('' if c is None else str(c) for c in r[:3]))
        else:
            with open(path, encoding='utf-8-sig', errors='ignore') as f:
                lines = f.read().splitlines()
        self.box.setPlainText('\n'.join(lines))


class PayeesTab(QWidget):
    changed = Signal()

    def __init__(self, ctx):
        super().__init__()
        self.ctx, self.db = ctx, ctx.db
        self.table = DataTable([
            (L('الاسم', 'Name'), 'text'), (L('الاسم بالإنجليزية', 'English name'), 'text'), (L('الهاتف', 'Phone'), 'text'),
            (L('عدد الشيكات', 'Cheques'), 'num'), (L('الحالة', 'Status'), 'select'), (L('ملاحظات', 'Notes'), 'text'),
        ], actions=[
            (L('تعديل', 'Edit'), lambda ids: self.edit(ids[0]), False, 'edit'),
            (L('تفعيل', 'Activate'), lambda ids: self._active(ids, True), False, 'on'),
            (L('إيقاف', 'Deactivate'), lambda ids: self._active(ids, False), False, 'off'),
            (L('حذف', 'Delete'), self._delete, True, 'trash'),
        ], export_name='payees', title=L('المستفيدون', 'Payees'))
        self.table.activated.connect(self.edit)
        add = button(L('مستفيد جديد', 'New payee'), 'plus', 'primary')
        add.clicked.connect(lambda: self.edit(None))
        imp = button(L('إضافة قائمة / استيراد', 'Add a list / import'), 'import',
                     tip=L('لصق قائمة أسماء أو اختيار ملف Excel / CSV', 'Paste a list of names or choose an Excel / CSV file'))
        imp.clicked.connect(self.import_list)
        top = QHBoxLayout()
        top.setSpacing(8)
        top.addWidget(add)
        top.addWidget(imp)
        top.addStretch(1)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(10)
        lay.addLayout(top)
        lay.addWidget(self.table)
        self.reload()

    def reload(self):
        self.table.set_rows([(p['id'], [p['name'], p.get('name_en') or '', p.get('phone') or '',
                                        (p['cheques'], p['cheques']),
                                        L('فعّال', 'Active') if p['is_active'] else L('موقوف', 'Inactive'),
                                        p.get('notes') or '']) for p in self.db.payees()])

    def edit(self, pid):
        cur = next((p for p in self.db.payees() if p['id'] == pid), None) if pid else None
        dlg = PayeeDialog(self, cur)
        if dlg.exec():
            v = dlg.values()
            if not v['name']:
                error(self, L('أدخل الاسم.', 'Enter the name.'))
                return
            try:
                self.db.save_payee(v, pid)
            except Exception as e:  # noqa: BLE001 — duplicate name
                error(self, L('يوجد مستفيد بالاسم نفسه.', 'A payee with this name exists.') + f'\n{e}')
                return
            self.reload()
            self.changed.emit()

    def _active(self, ids, on):
        self.db.set_active('payees', ids, on)
        self.reload()
        self.changed.emit()

    def _delete(self, ids):
        used = {p['id']: p['name'] for p in self.db.payees() if p['cheques']}
        if not confirm(self, L(f'حذف {len(ids)} من المستفيدين؟ من له شيكات يُترك (أوقفه بدلاً من ذلك).',
                               f'Delete {len(ids)} payee(s)? Those with cheques are kept (deactivate them instead).')):
            return
        done, kept = self.db.delete_rows('payees', ids, guard=lambda i: used.get(i))
        msg = L(f'حُذف: {done}', f'Deleted: {done}')
        if kept:
            msg += '\n' + L('تُرك (له شيكات): ', 'Kept (has cheques): ') + '، '.join(kept)
        info(self, msg)
        self.reload()
        self.changed.emit()

    def import_list(self):
        dlg = ImportPayeesDialog(self)
        if not dlg.exec():
            return
        box = dlg.box
        existing = {p['name'] for p in self.db.payees()}
        added = skipped = 0
        for line in box.toPlainText().splitlines():
            parts = [x.strip() for x in (line.split('\t') if '\t' in line else line.split(','))]
            name = parts[0] if parts else ''
            if not name or name in ('الاسم', 'name', 'Name'):
                continue
            if name in existing:
                skipped += 1
                continue
            self.db.save_payee({'name': name, 'name_en': parts[1] if len(parts) > 1 and parts[1] else None,
                                'phone': parts[2] if len(parts) > 2 and parts[2] else None})
            existing.add(name)
            added += 1
        info(self, L(f'أُضيف {added}، وتُرك {skipped} (موجود مسبقاً).', f'Added {added}, skipped {skipped} (already there).'))
        self.reload()
        self.changed.emit()


# ───────────────────────────── banks & accounts ─────────────────────────────
class AccountDialog(QDialog):
    def __init__(self, parent, db, acc=None):
        super().__init__(parent)
        self.setWindowTitle(L('حساب بنكي', 'Bank account'))
        acc = acc or {}
        self.bank = QComboBox()
        for b in db.banks():
            self.bank.addItem(b['name'], b['id'])
        if acc.get('bank_id'):
            self.bank.setCurrentIndex(max(0, self.bank.findData(acc['bank_id'])))
        self.name = QLineEdit(acc.get('name') or '')
        self.number = QLineEdit(acc.get('number') or '')
        self.currency = QComboBox()
        for k in CURRENCIES:
            self.currency.addItem(k, k)
        self.currency.setCurrentIndex(max(0, self.currency.findData(acc.get('currency') or 'JOD')))
        self.next_no = QSpinBox()
        self.next_no.setRange(1, 999999999)
        self.next_no.setValue(int(acc.get('next_no') or 1))
        self.last_no = QSpinBox()
        self.last_no.setRange(0, 999999999)
        self.last_no.setSpecialValueText(L('غير محدد', 'Not set'))
        self.last_no.setValue(int(acc.get('last_no') or 0))
        self.notes = QLineEdit(acc.get('notes') or '')
        self.active = a11y(QCheckBox(L('فعّال', 'Active')), L('الحساب فعّال', 'Account active'))
        self.active.setChecked(bool(acc.get('is_active', 1)))
        self.bank.setMinimumWidth(260)
        ltr(self.number, 220)
        fit(self.currency, 120)
        for sp in (self.next_no, self.last_no):
            ltr(sp, 160)
            sp.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        g = form_grid()
        field(g, 0, 0, L('البنك (القالب)', 'Bank (template)'), self.bank)
        field(g, 1, 0, L('اسم الحساب *', 'Account name *'), self.name)
        field(g, 2, 0, L('رقم الحساب', 'Account number'), self.number)
        field(g, 3, 0, L('العملة', 'Currency'), self.currency)
        field(g, 4, 0, L('رقم الشيك التالي', 'Next cheque number'), self.next_no)
        field(g, 5, 0, L('آخر رقم في دفتر الشيكات', 'Last number in the cheque book'), self.last_no)
        field(g, 6, 0, L('ملاحظات', 'Notes'), self.notes)
        g.addWidget(self.active, 7, 1)
        _dialog_layout(self, g)

    def values(self):
        return {'bank_id': self.bank.currentData(), 'name': self.name.text().strip(),
                'number': self.number.text().strip() or None, 'currency': self.currency.currentData(),
                'next_no': self.next_no.value(), 'last_no': self.last_no.value() or None,
                'notes': self.notes.text().strip() or None, 'is_active': 1 if self.active.isChecked() else 0}


class BanksTab(QWidget):
    changed = Signal()

    def __init__(self, ctx):
        super().__init__()
        self.ctx, self.db = ctx, ctx.db
        self.banks = DataTable([(L('البنك / القالب', 'Bank / template'), 'text'), (L('بالإنجليزية', 'English'), 'text'),
                                (L('المقاس (مم)', 'Size (mm)'), 'text'), (L('الحسابات', 'Accounts'), 'num'),
                                (L('الحالة', 'Status'), 'select')],
                               actions=[(L('معايرة القالب', 'Calibrate template'), lambda ids: self.calibrate(ids[0]), False, 'calibrate'),
                                        (L('نسخ القالب', 'Duplicate'), lambda ids: self.duplicate(ids[0]), False, 'copy'),
                                        (L('تفعيل', 'Activate'), lambda ids: self._active('banks', ids, True), False, 'on'),
                                        (L('إيقاف', 'Deactivate'), lambda ids: self._active('banks', ids, False), False, 'off'),
                                        (L('حذف', 'Delete'), self._delete_banks, True, 'trash')], export_name='banks',
                               title=L('البنوك وقوالب الشيكات', 'Banks and cheque templates'))
        self.banks.activated.connect(self.calibrate)
        self.accounts = DataTable([(L('الحساب', 'Account'), 'text'), (L('البنك', 'Bank'), 'select'),
                                   (L('رقم الحساب', 'Number'), 'text'), (L('العملة', 'Currency'), 'select'),
                                   (L('الشيك التالي', 'Next cheque'), 'num'), (L('آخر رقم في الدفتر', 'Last in book'), 'num'),
                                   (L('الشيكات', 'Cheques'), 'num'), (L('الحالة', 'Status'), 'select')],
                                  actions=[(L('تعديل', 'Edit'), lambda ids: self.edit_account(ids[0]), False, 'edit'),
                                           (L('جعله الافتراضي', 'Make default'), self._default, False, 'star'),
                                           (L('تفعيل', 'Activate'), lambda ids: self._active('accounts', ids, True), False, 'on'),
                                           (L('إيقاف', 'Deactivate'), lambda ids: self._active('accounts', ids, False), False, 'off'),
                                           (L('حذف', 'Delete'), self._delete_accounts, True, 'trash')], export_name='accounts',
                                  title=L('الحسابات البنكية', 'Bank accounts'))
        self.accounts.activated.connect(self.edit_account)
        add_bank = button(L('بنك / قالب جديد', 'New bank / template'), 'bank',
                          tip=L('إضافة بنك ومعايرة قالب شيكاته', 'Add a bank and calibrate its cheque template'))
        add_bank.clicked.connect(self.add_bank)
        add_acc = button(L('حساب بنكي جديد', 'New bank account'), 'plus', 'primary')
        add_acc.clicked.connect(lambda: self.edit_account(None))
        gb1 = card(L('البنوك وقوالب الشيكات', 'Banks and cheque templates'))
        gb2 = card(L('الحسابات البنكية ودفاتر الشيكات', 'Bank accounts and cheque books'))
        l1, l2 = QVBoxLayout(gb1), QVBoxLayout(gb2)
        h1, h2 = QHBoxLayout(), QHBoxLayout()
        h1.addWidget(add_bank)
        h1.addStretch(1)
        h2.addWidget(add_acc)
        h2.addStretch(1)
        l1.addLayout(h1)
        l1.addWidget(self.banks)
        l2.addLayout(h2)
        l2.addWidget(self.accounts)
        sp = QSplitter(Qt.Vertical)
        sp.setChildrenCollapsible(False)
        sp.setHandleWidth(10)
        sp.addWidget(gb2)
        sp.addWidget(gb1)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(14, 0, 14, 14)
        lay.addWidget(sp)
        self.reload()

    def reload(self):
        counts = {}
        for a in self.db.accounts():
            counts[a['bank_id']] = counts.get(a['bank_id'], 0) + 1
        rows = []
        for b in self.db.banks():
            t = self.db.template(b['id'])
            rows.append((b['id'], [b['name'], b.get('name_en') or '', f'{t["width_mm"]:g} × {t["height_mm"]:g}',
                                   (counts.get(b['id'], 0), counts.get(b['id'], 0)),
                                   L('فعّال', 'Active') if b['is_active'] else L('موقوف', 'Inactive')]))
        self.banks.set_rows(rows)
        default = self.db.get('default_account')
        rows = []
        for a in self.db.accounts():
            n = self.db.account_in_use(a['id'])
            name = a['name'] + (L(' (الافتراضي)', ' (default)') if str(a['id']) == default else '')
            rows.append((a['id'], [name, a['bank_name'], a.get('number') or '', a['currency'],
                                   (a['next_no'], a['next_no']), (a.get('last_no') or '', a.get('last_no') or 0),
                                   (n, n), L('فعّال', 'Active') if a['is_active'] else L('موقوف', 'Inactive')]))
        self.accounts.set_rows(rows)

    def add_bank(self):
        from .template_editor import TemplateEditor
        dlg = TemplateEditor(self, self.ctx, None)
        if dlg.exec():
            self.reload()
            self.changed.emit()

    def calibrate(self, bank_id):
        from .template_editor import TemplateEditor
        dlg = TemplateEditor(self, self.ctx, bank_id)
        if dlg.exec():
            self.reload()
            self.changed.emit()

    def duplicate(self, bank_id):
        b = self.db.bank(bank_id)
        t = self.db.template(bank_id)
        self.db.save_bank({'name': b['name'] + L(' (نسخة)', ' (copy)'), 'name_en': b.get('name_en'), 'template': t,
                           'background': b.get('background')})
        self.reload()
        self.changed.emit()

    def edit_account(self, acc_id):
        acc = self.db.account(acc_id) if acc_id else None
        dlg = AccountDialog(self, self.db, acc)
        if dlg.exec():
            v = dlg.values()
            if not v['name'] or not v['bank_id']:
                error(self, L('أدخل اسم الحساب واختر البنك.', 'Enter the account name and choose the bank.'))
                return
            new_id = self.db.save_account(v, acc_id)
            if not self.db.get('default_account'):
                self.db.set('default_account', new_id)
            self.reload()
            self.changed.emit()

    def _default(self, ids):
        self.db.set('default_account', ids[0])
        self.reload()
        self.changed.emit()

    def _active(self, table, ids, on):
        self.db.set_active(table, ids, on)
        self.reload()
        self.changed.emit()

    def _delete_banks(self, ids):
        if not confirm(self, L(f'حذف {len(ids)} من البنوك/القوالب؟ البنك الذي له حسابات يُترك.',
                               f'Delete {len(ids)} bank(s)? Banks with accounts are kept.')):
            return
        names = {b['id']: b['name'] for b in self.db.banks()}
        done, kept = self.db.delete_rows('banks', ids, guard=lambda i: names[i] if self.db.bank_in_use(i) else None)
        info(self, L(f'حُذف: {done}', f'Deleted: {done}') + (('\n' + L('تُرك (له حسابات): ', 'Kept (has accounts): ') + '، '.join(kept)) if kept else ''))
        self.reload()
        self.changed.emit()

    def _delete_accounts(self, ids):
        if not confirm(self, L(f'حذف {len(ids)} من الحسابات؟ الحساب الذي له شيكات يُترك (أوقفه بدلاً من ذلك).',
                               f'Delete {len(ids)} account(s)? Accounts with cheques are kept (deactivate them instead).')):
            return
        names = {a['id']: a['name'] for a in self.db.accounts()}
        done, kept = self.db.delete_rows('accounts', ids, guard=lambda i: names[i] if self.db.account_in_use(i) else None)
        info(self, L(f'حُذف: {done}', f'Deleted: {done}') + (('\n' + L('تُرك (له شيكات): ', 'Kept (has cheques): ') + '، '.join(kept)) if kept else ''))
        self.reload()
        self.changed.emit()
